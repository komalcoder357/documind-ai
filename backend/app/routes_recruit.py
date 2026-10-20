import csv, io, os, uuid
from datetime import datetime, timezone
from typing import Literal
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import Response
from fastapi.encoders import jsonable_encoder
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session
import docx
from . import config, cv, extract, jobs
from .db import get_db
from .deps import current_user, ws_ids, primary_ws, audit
from .models import *
from .routes_docs import job_out

router = APIRouter(prefix="/api/v1/recruitment")
DISCLAIMER = ("This is a job-requirement match score based on documented rules and available evidence. It is not a measure of a person's worth or guaranteed job performance. "
              "Missing evidence is not evidence of absence. All results require meaningful human review; no candidate is automatically accepted or rejected.")

class CriterionIn(BaseModel):
    type: Literal["skill", "keyword", "experience", "education"]; level: Literal["required", "preferred"] = "required"
    target: str | int; aliases: list[str] = Field(default_factory=list, max_length=10); weight: float = Field(gt=0, le=10)
class ProjectIn(BaseModel):
    title: str = Field(min_length=1, max_length=200); job_title: str = Field(min_length=1, max_length=200); job_description: str = Field("", max_length=20000)
    required_skills: list[str] = Field(default_factory=list, max_length=30); preferred_skills: list[str] = Field(default_factory=list, max_length=30)
    min_years: int | None = Field(None, ge=0, le=50); education: Literal["diploma", "bachelor", "master", "phd"] | None = None
    criteria: list[CriterionIn] | None = None
class AnalyzeIn(BaseModel): check_portfolio: bool = True

def finalize(criteria):
    out = []
    for i, c in enumerate(criteria, 1):
        d = c.model_dump() if hasattr(c, "model_dump") else dict(c)
        if d["type"] == "experience":
            try: d["target"] = int(d["target"])
            except (TypeError, ValueError): raise HTTPException(422, "Experience target must be a number of years.")
        d["id"] = i; d["name"] = d["target"] if d["type"] in ("skill", "keyword") else f"{d['type'].title()}: {d['target']}"; out.append(d)
    return out

def proj_or_404(db, user, pid):
    p = db.get(RecruitProject, pid)
    if not p or p.workspace_id not in ws_ids(db, user): raise HTTPException(404, "Project not found.")
    return p
def cand_or_404(db, user, cid):
    c = db.get(Candidate, cid)
    if not c or c.workspace_id not in ws_ids(db, user): raise HTTPException(404, "Candidate not found.")
    return c
def pout(p): return {"id": p.id, "title": p.title, "job_title": p.job_title, "job_description": p.job_description, "criteria": p.criteria, "created_at": p.created_at}

@router.post("/projects", status_code=201)
def create_project(body: ProjectIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    crit = finalize(body.criteria) if body.criteria else cv.build_criteria(body.required_skills, body.preferred_skills, body.min_years, body.education)
    p = RecruitProject(workspace_id=primary_ws(db, user), owner_id=user.id, title=body.title, job_title=body.job_title, job_description=body.job_description, criteria=crit)
    db.add(p); db.commit(); audit(db, user.id, "project.create"); return pout(p)

@router.get("/projects")
def list_projects(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return [pout(p) for p in db.scalars(select(RecruitProject).where(RecruitProject.workspace_id.in_(ws_ids(db, user))).order_by(RecruitProject.id.desc()))]

@router.get("/projects/{pid}")
def get_project(pid: int, user: User = Depends(current_user), db: Session = Depends(get_db)): return pout(proj_or_404(db, user, pid))

@router.post("/projects/{pid}/criteria")
def set_criteria(pid: int, body: list[CriterionIn], user: User = Depends(current_user), db: Session = Depends(get_db)):
    p = proj_or_404(db, user, pid)
    if not body: raise HTTPException(422, "At least one criterion is required.")
    p.criteria = finalize(body); db.commit(); return {"criteria": p.criteria, "formula": cv.FORMULA}

@router.delete("/projects/{pid}")
def delete_project(pid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    p = proj_or_404(db, user, pid)
    for c in db.scalars(select(Candidate).where(Candidate.project_id == p.id)):
        try: os.remove(os.path.join(config.STORAGE_DIR, c.storage_key))
        except FileNotFoundError: pass
    db.delete(p); db.commit(); return {"ok": True}

@router.post("/projects/{pid}/candidates", status_code=201)
async def upload_cvs(pid: int, files: list[UploadFile] = File(...), user: User = Depends(current_user), db: Session = Depends(get_db)):
    p = proj_or_404(db, user, pid)
    if len(files) > 30: raise HTTPException(413, "At most 30 CVs per upload.")
    out = []
    for f in files:
        name = os.path.basename(f.filename or "")[:255]
        try:
            raw = await f.read(config.MAX_UPLOAD + 1)
            if len(raw) > config.MAX_UPLOAD: raise extract.ExtractError("File too large.", 413)
            kind = extract.sniff(raw, name); pages, warns = extract.extract_pages(raw, kind, config.MAX_PAGES)
            text = "\n".join(t for _, t in pages)
            if len(text.strip()) < 20: raise extract.ExtractError("No readable text found in this CV.")
        except extract.ExtractError as e: out.append({"filename": name, "status": "rejected", "error": str(e)}); continue
        key = uuid.uuid4().hex
        os.makedirs(config.STORAGE_DIR, exist_ok=True)
        with open(os.path.join(config.STORAGE_DIR, key), "wb") as fh: fh.write(raw)
        prof = cv.parse_cv(text); c = Candidate(project_id=p.id, workspace_id=p.workspace_id, filename=name, storage_key=key, text=text, profile=prof, warnings=warns + prof["extraction_warnings"])
        db.add(c); db.flush(); out.append({"id": c.id, "filename": name, "status": "uploaded", "warnings": c.warnings})
    db.commit(); audit(db, user.id, "cv.upload"); return {"results": out}

@router.post("/projects/{pid}/analyze")
def analyze(pid: int, body: AnalyzeIn = AnalyzeIn(), user: User = Depends(current_user), db: Session = Depends(get_db)):
    p = proj_or_404(db, user, pid)
    if not p.criteria: raise HTTPException(422, "Define and save job criteria before analysis.")
    cands = list(db.scalars(select(Candidate).where(Candidate.project_id == p.id)))
    if not cands: raise HTTPException(422, "Upload at least one CV first.")
    js = []
    for c in cands:
        c.profile = {**c.profile, "_check_portfolio": body.check_portfolio}; c.status = "queued"; db.commit()
        js.append(job_out(jobs.enqueue(db, "candidate", c.id, p.workspace_id)))
    audit(db, user.id, "project.analyze"); return {"jobs": js}

def cand_summary(c, a):
    prof = {k: v for k, v in c.profile.items() if not k.startswith("_")}
    return {"id": c.id, "filename": c.filename, "status": c.status, "error": c.error, "name": prof.get("name"), "overall": a.overall if a else None, "coverage": a.coverage if a else None,
            "criteria": [{k: r[k] for k in ("id", "name", "type", "level", "weight", "status", "value", "explanation")} | {"evidence_count": len(r["evidence"])} for r in (a.criteria if a else [])],
            "portfolio": [{"url": l["url"], "kind": l["kind"], "status": l["status"], "reason": l.get("reason"), "repos": [r["name"] for r in l.get("repos", []) if r.get("status") == "checked"],
                           "repo_details": [{k: r.get(k) for k in ("name", "has_readme", "has_tests", "has_docs", "fork")} for r in l.get("repos", []) if r.get("status") == "checked"]} for l in c.portfolio],
            "warnings": c.warnings}

def results_payload(db, p, sort="overall", status=None):
    rows = [(c, db.scalar(select(Assessment).where(Assessment.candidate_id == c.id))) for c in db.scalars(select(Candidate).where(Candidate.project_id == p.id))]
    if status: rows = [r for r in rows if r[0].status == status]
    rows.sort(key=(lambda r: r[0].filename.lower()) if sort == "filename" else (lambda r: -(r[1].overall if r[1] and r[1].overall is not None else -1)))
    return {"project": pout(p), "formula": cv.FORMULA, "disclaimer": DISCLAIMER, "candidates": [cand_summary(c, a) for c, a in rows]}

@router.get("/projects/{pid}/results")
def results(pid: int, sort: Literal["overall", "filename"] = "overall", status: str | None = None, user: User = Depends(current_user), db: Session = Depends(get_db)):
    return results_payload(db, proj_or_404(db, user, pid), sort, status)

@router.get("/candidates/{cid}/report")
def report(cid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    c = cand_or_404(db, user, cid); a = db.scalar(select(Assessment).where(Assessment.candidate_id == c.id)); p = db.get(RecruitProject, c.project_id)
    if not a: raise HTTPException(409, "This candidate has not been analyzed yet.")
    prof = {k: v for k, v in c.profile.items() if not k.startswith("_")}
    return {"overview": {"filename": c.filename, **{k: prof.get(k) for k in ("name", "email", "years_experience", "education_level", "links")}},
            "job": {"title": p.job_title, "criteria": p.criteria}, "overall": a.overall, "coverage": a.coverage, "formula": cv.FORMULA,
            "matched": [r for r in a.criteria if r["status"] in ("listed_in_cv", "supported_by_work_sample")],
            "partial": [r for r in a.criteria if r["status"] == "partial"],
            "missing_or_uncertain": [r for r in a.criteria if r["status"] in ("not_found", "unable_to_verify")],
            "criteria": a.criteria, "portfolio": [{k: v for k, v in l.items()} | {"repos": [{k: v for k, v in r.items() if k != "readme_text"} for r in l.get("repos", [])]} for l in c.portfolio],
            "interview_questions": a.interview_questions, "limitations": [DISCLAIMER, "Repository metadata does not prove authorship or skill level.", "Only text extracted from the CV and accessible public sources was considered."],
            "evaluated_at": a.created_at.isoformat() if a.created_at else None}

@router.post("/candidates/{cid}/interview-questions")
def questions(cid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    c = cand_or_404(db, user, cid); a = db.scalar(select(Assessment).where(Assessment.candidate_id == c.id))
    if not a: raise HTTPException(409, "Analyze the candidate first.")
    a.interview_questions = cv.interview_questions(c.profile, a.criteria); db.commit()
    return {"questions": a.interview_questions, "notice": "Generated questions do not prove ability; they structure verification at interview."}

def render_csv(payload):
    crit = [c["name"] for c in payload["project"]["criteria"]]; buf = io.StringIO(); w = csv.writer(buf)
    w.writerow(["candidate_file", "overall_match_score", "evidence_coverage"] + crit)
    for c in payload["candidates"]:
        by = {r["id"]: r for r in c["criteria"]}
        w.writerow([c["filename"], "" if c["overall"] is None else c["overall"], "" if c["coverage"] is None else c["coverage"]] + [by.get(x["id"], {}).get("status", "") for x in payload["project"]["criteria"]])
    return buf.getvalue()

def render_docx(title, payload):
    d = docx.Document(); d.add_heading(title, 0)
    if "candidates" in payload:
        d.add_paragraph(payload.get("formula", "")); crit = payload["project"]["criteria"]
        t = d.add_table(rows=1, cols=3 + len(crit)); t.style = "Table Grid"
        for i, h in enumerate(["Candidate file", "Match score", "Coverage"] + [c["name"] for c in crit]): t.rows[0].cells[i].text = str(h)
        for c in payload["candidates"]:
            by = {r["id"]: r for r in c["criteria"]}; r = t.add_row().cells
            r[0].text, r[1].text, r[2].text = c["filename"], str(c["overall"] if c["overall"] is not None else "n/a"), str(c["coverage"] if c["coverage"] is not None else "n/a")
            for i, x in enumerate(crit): r[3 + i].text = by.get(x["id"], {}).get("status", "")
        d.add_heading("Limitations and human review", 1); d.add_paragraph(payload.get("disclaimer", ""))
    else:
        import json; d.add_paragraph(json.dumps(payload, indent=2, default=str)[:30000])
    b = io.BytesIO(); d.save(b); return b.getvalue()

@router.post("/projects/{pid}/export", status_code=201)
def export(pid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    p = proj_or_404(db, user, pid); payload = jsonable_encoder(results_payload(db, p))
    r = Report(workspace_id=p.workspace_id, owner_id=user.id, kind="candidate_comparison", title=f"Comparison: {p.title}", payload=payload); db.add(r); db.commit()
    audit(db, user.id, "report.create"); return {"report_id": r.id}

# ---- reports (shared) ----
rep = APIRouter(prefix="/api/v1/reports")
class SaveIn(BaseModel): title: str = Field(min_length=1, max_length=200); kind: str = Field(max_length=30); payload: dict
def rep_or_404(db, user, rid):
    r = db.get(Report, rid)
    if not r or r.workspace_id not in ws_ids(db, user): raise HTTPException(404, "Report not found.")
    return r
@rep.get("")
def list_reports(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return [{"id": r.id, "kind": r.kind, "title": r.title, "created_at": r.created_at} for r in db.scalars(select(Report).where(Report.workspace_id.in_(ws_ids(db, user))).order_by(Report.id.desc()).limit(100))]
@rep.post("", status_code=201)
def save(body: SaveIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    r = Report(workspace_id=primary_ws(db, user), owner_id=user.id, kind=body.kind, title=body.title, payload=jsonable_encoder(body.payload)); db.add(r); db.commit(); return {"report_id": r.id}
@rep.get("/{rid}")
def get_report(rid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    r = rep_or_404(db, user, rid); return {"id": r.id, "kind": r.kind, "title": r.title, "payload": r.payload, "created_at": r.created_at}
@rep.get("/{rid}/download")
def dl(rid: int, format: Literal["csv", "docx", "json"] = "docx", user: User = Depends(current_user), db: Session = Depends(get_db)):
    r = rep_or_404(db, user, rid); safe = f"report_{r.id}"
    if format == "csv":
        if "candidates" not in r.payload: raise HTTPException(422, "CSV is only available for candidate comparisons.")
        return Response(render_csv(r.payload), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{safe}.csv"'})
    if format == "json":
        import json; return Response(json.dumps(r.payload, default=str), media_type="application/json", headers={"Content-Disposition": f'attachment; filename="{safe}.json"'})
    return Response(render_docx(r.title, r.payload), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="{safe}.docx"'})
@rep.delete("/{rid}")
def del_report(rid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    db.delete(rep_or_404(db, user, rid)); db.commit(); return {"ok": True}
